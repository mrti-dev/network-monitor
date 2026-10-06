from copy import deepcopy
from pathlib import Path
from sys import argv

from docx import Document
from docx.oxml.ns import qn


TEXT = {
    1: "THIẾT KẾ HỆ THỐNG CÁ NHÂN",
    2: "Họ và tên : Vũ Mạnh Tiến",
    3: "MSV :",
    6: "PHẦN 1: THIẾT KẾ KỸ THUẬT TỪ SRS NHÓM THỐNG NHẤT",
    7: "1.1. Xác định module được phân công",
    8: "Tên Module: M2 – Quản lý Phòng ban và Hồ sơ nhân sự.",
    9: "Phạm vi nghiệp vụ: CRUD phòng ban, chức vụ, chỉ định Trưởng phòng; CRUD hồ sơ nhân viên; điều chuyển phòng ban và bổ nhiệm chức vụ (P02); tìm kiếm, lọc và phân trang bằng Ajax không tải lại trang (R07).",
    10: "1.2. Đặc tả chi tiết 3 Use Case chính",
    11: "Use Case 1: UC-EMP-01 - Thêm mới hồ sơ nhân viên",
    12: "Tác nhân: Chuyên viên HR.",
    13: "Mô tả: Cho phép HR tạo hồ sơ nhân viên mới, nhập thông tin định danh và gán phòng ban, chức vụ, trạng thái làm việc ban đầu.",
    14: "Tiền điều kiện: HR đã đăng nhập; danh mục phòng ban và chức vụ đã có dữ liệu hợp lệ.",
    15: "Luồng sự kiện chính (Main Flow):",
    16: "HR chọn chức năng \"Thêm mới nhân viên\".",
    17: "Hệ thống hiển thị biểu mẫu tạo hồ sơ nhân viên.",
    18: "HR nhập họ tên, ngày sinh, giới tính, CCCD, email, số điện thoại, địa chỉ và ngày vào làm.",
    19: "HR chọn phòng ban, chức vụ và trạng thái làm việc ban đầu.",
    20: "Hệ thống kiểm tra dữ liệu bắt buộc, định dạng và tính duy nhất của CCCD, email.",
    21: "Hệ thống sinh EmployeeCode và lưu Entity Employee vào CSDL.",
    22: "Hệ thống chuyển đến trang chi tiết và thông báo thêm hồ sơ thành công.",
    23: "Luồng ngoại lệ (Exception Flow):",
    24: "Bước 5: Nếu CCCD hoặc email đã tồn tại, hệ thống hiển thị lỗi tại trường nhập và không lưu dữ liệu.",
    25: "Hậu điều kiện: Một hồ sơ Employee mới được lưu và liên kết với Department, Position đã chọn.",
    26: "",
    27: "Use Case 2: UC-EMP-02 - Xem danh sách, tìm kiếm và lọc hồ sơ nhân viên",
    28: "Tác nhân: Chuyên viên HR, Trưởng phòng (Manager).",
    29: "Mô tả: Cho phép người dùng có quyền tra cứu hồ sơ theo mã, tên, phòng ban và trạng thái; kết quả được phân trang bằng Ajax.",
    30: "Tiền điều kiện: Người dùng đã đăng nhập và được cấp quyền xem danh sách hồ sơ theo phạm vi quản lý.",
    31: "Luồng sự kiện chính (Main Flow):",
    32: "Người dùng mở màn hình danh sách nhân viên.",
    33: "Hệ thống tải trang dữ liệu đầu tiên theo quyền truy cập.",
    34: "Người dùng nhập từ khóa hoặc chọn Department, WorkStatus.",
    35: "Ajax gửi Keyword, DepartmentId, WorkStatus và PageNumber đến Controller.",
    36: "Hệ thống truy vấn, lọc, phân trang và trả về danh sách phù hợp.",
    37: "Giao diện cập nhật vùng kết quả mà không tải lại toàn trang.",
    38: "Luồng ngoại lệ: Nếu không có dữ liệu phù hợp hoặc yêu cầu Ajax lỗi, hệ thống hiển thị thông báo và giữ nguyên bộ lọc.",
    39: "PHẦN 2: THIẾT KẾ DỮ LIỆU",
    40: "2.1. Danh sách Entity/Domain Model",
    41: "Lớp Employee (Hồ sơ nhân viên):",
    42: "EmployeeId (int, PK): Khóa chính.",
    43: "EmployeeCode (string, Unique): Mã nhân viên.",
    44: "FullName (string), IdentityCard (string, Unique), Email (string, Unique), PhoneNumber (string).",
    45: "DepartmentId (int, FK), PositionId (int, FK): Cơ cấu công tác hiện tại.",
    46: "WorkStatus (int), HireDate (DateTime): Trạng thái và ngày vào làm.",
    47: "IsDeleted (bool): Cờ xóa mềm bảo lưu lịch sử.",
    48: "Lớp Department (Phòng ban):",
    49: "DepartmentId (int, PK): Khóa chính.",
    50: "DepartmentCode (string, Unique), DepartmentName (string): Mã và tên phòng ban.",
    51: "ManagerId (int, FK, nullable), PhoneNumber (string), IsDeleted (bool).",
    52: "Lớp Position (Chức vụ):",
    53: "PositionId (int, PK): Khóa chính.",
    54: "PositionName (string), BaseSalaryRate (decimal): Tên chức vụ và hệ số lương.",
    55: "Allowance (decimal), IsDeleted (bool): Phụ cấp và trạng thái xóa mềm.",
    56: "Lớp JobTransferHistory (Lịch sử điều chuyển):",
    57: "TransferId (int, PK), EmployeeId (int, FK), OldDepartmentId, NewDepartmentId, OldPositionId, NewPositionId, TransferDate, Reason.",
    58: "2.2. Danh sách ViewModel / EditModel",
    59: "EmployeeCreateViewModel: Model nhập hồ sơ mới gồm FullName, DateOfBirth, Gender, IdentityCard, Email, PhoneNumber, Address, AvatarUrl, DepartmentId, PositionId, WorkStatus, HireDate; sử dụng Data Annotations để kiểm tra dữ liệu.",
    60: "EmployeeFilterViewModel: Model tìm kiếm và phân trang gồm Keyword, DepartmentId, WorkStatus, PageNumber, PageSize; dùng cho yêu cầu Ajax.",
    61: "JobTransferEditModel: Model điều chuyển gồm EmployeeId, NewDepartmentId, NewPositionId, TransferDate, Reason; kiểm tra phòng ban mới và ràng buộc Trưởng phòng.",
    62: "PHẦN 3: THIẾT KẾ KIẾN TRÚC LUỒNG MVC CRUD FLOW",
    63: "Luồng MVC cho chức năng: Điều chuyển phòng ban và bổ nhiệm chức vụ mới",
    64: "Client Request (User Action): HR chọn \"Điều chuyển\" tại hồ sơ nhân viên. Trình duyệt gửi HTTP GET request kèm EmployeeId lên máy chủ.",
    65: "Controller (GET):",
    66: "EmployeeController nhận request vào Action Transfer(int id).",
    67: "Truy vấn CSDL qua Entity Framework / Repository để lấy Employee cùng Department và Position hiện tại.",
    68: "Khởi tạo JobTransferEditModel với EmployeeId và thông tin cơ cấu hiện tại.",
    69: "Controller trả ViewModel về cho View.",
    70: "View (Hiển thị & Form POST): View Transfer.cshtml hiển thị phòng ban, chức vụ hiện tại và danh mục mới. HR nhập ngày hiệu lực, lý do rồi gửi HTTP POST request.",
    71: "Controller (POST - Xử lý Logic):",
    72: "EmployeeController nhận request vào Action Transfer(JobTransferEditModel model).",
    73: "Validation: Kiểm tra ModelState.IsValid, phòng ban mới khác phòng ban hiện tại và ràng buộc thay thế Trưởng phòng.",
    74: "Nếu dữ liệu sai: Trả lại View kèm thông báo lỗi.",
    75: "Nếu dữ liệu đúng: Cập nhật DepartmentId, PositionId của Employee và tạo Entity JobTransferHistory lưu giá trị cũ, mới.",
    76: "Model / Data Access (Lưu trữ):",
    77: "Entity Framework thực hiện cập nhật trong một transaction và gọi _dbContext.SaveChanges().",
    78: "Câu lệnh SQL UPDATE Employee và INSERT JobTransferHistory được thực thi xuống SQL Server.",
    79: "Response (Phản hồi): Controller RedirectToAction(\"Details\"), trả về hồ sơ đã cập nhật và thông báo \"Điều chuyển nhân sự thành công\".",
}


THIRD_USE_CASE = [
    (27, "Use Case 3: UC-ORG-04 - Điều chuyển phòng ban và bổ nhiệm chức vụ mới"),
    (28, "Tác nhân: Chuyên viên HR."),
    (29, "Mô tả: Cho phép HR cập nhật phòng ban, chức vụ của nhân viên và lưu lịch sử điều chuyển phục vụ truy vết."),
    (30, "Tiền điều kiện: HR đã đăng nhập; nhân viên đang làm việc; Department và Position mới tồn tại."),
    (31, "Luồng sự kiện chính (Main Flow):"),
    (32, "HR tìm và mở hồ sơ nhân viên cần điều chuyển."),
    (33, "Hệ thống hiển thị Department và Position hiện tại."),
    (34, "HR chọn Department, Position mới, nhập TransferDate và Reason."),
    (35, "Hệ thống kiểm tra trạng thái nhân viên và ràng buộc Trưởng phòng."),
    (36, "Cập nhật Employee và tạo JobTransferHistory trong cùng transaction."),
    (37, "Hệ thống thông báo điều chuyển thành công."),
    (38, "Luồng ngoại lệ: Nếu phòng ban mới trùng phòng ban hiện tại hoặc chưa chỉ định người thay thế Trưởng phòng, hệ thống từ chối lưu."),
]


def replace_paragraph(paragraph, text):
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._element.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._element.rPr)
    for child in list(paragraph._element):
        if child.tag != qn("w:pPr"):
            paragraph._element.remove(child)
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._element.insert(0, run_properties)


def clone_with_text(source, text):
    clone = deepcopy(source._element)
    for child in list(clone):
        if child.tag != qn("w:pPr"):
            clone.remove(child)
    source_rpr = source.runs[0]._element.rPr if source.runs and source.runs[0]._element.rPr is not None else None
    run = clone.makeelement(qn("w:r"))
    if source_rpr is not None:
        run.append(deepcopy(source_rpr))
    text_node = clone.makeelement(qn("w:t"))
    text_node.text = text
    run.append(text_node)
    clone.append(run)
    return clone


def copy_paragraph_properties(source, target):
    current = target._element.find(qn("w:pPr"))
    if current is not None:
        target._element.remove(current)
    source_properties = source._element.find(qn("w:pPr"))
    if source_properties is not None:
        target._element.insert(0, deepcopy(source_properties))


def main(reference, output):
    doc = Document(reference)
    original = list(doc.paragraphs)
    for index, text in TEXT.items():
        replace_paragraph(original[index - 1], text)
    copy_paragraph_properties(original[51], original[55])
    original[25]._element.getparent().remove(original[25]._element)

    section_two = original[38]._element
    postcondition = clone_with_text(original[24], "Hậu điều kiện: Danh sách Employee phù hợp được hiển thị theo đúng quyền và điều kiện lọc.")
    section_two.addprevious(postcondition)
    for source_index, text in THIRD_USE_CASE:
        section_two.addprevious(clone_with_text(original[source_index - 1], text))
    section_two.addprevious(clone_with_text(original[24], "Hậu điều kiện: Employee được cập nhật và một JobTransferHistory mới được lưu."))

    doc.core_properties.title = "Thiết kế hệ thống cá nhân phân hệ phòng ban và hồ sơ nhân sự"
    doc.core_properties.author = "Vũ Mạnh Tiến"
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    main(argv[1], argv[2])
