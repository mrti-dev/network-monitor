from copy import deepcopy
from pathlib import Path
from sys import argv

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PARAGRAPHS = {
    1: "BÁO CÁO CÁ NHÂN",
    2: "THIẾT KẾ HỆ THỐNG",
    3: "Đề tài: Quản lý nhân sự trong doanh nghiệp nhỏ\nSinh viên: Vũ Mạnh Tiến\nModule/phân hệ: Quản lý phòng ban và hồ sơ nhân sự",
    4: "1. GIỚI THIỆU PHẦN CÁ NHÂN",
    5: "Theo bảng phân công của nhóm, em phụ trách phân hệ Phòng ban và Hồ sơ nhân sự, gồm quản lý phòng ban, chức vụ, chỉ định Trưởng phòng, quản lý hồ sơ nhân viên và quy trình điều chuyển. Em đồng thời triển khai tìm kiếm, lọc và phân trang bằng Ajax để cập nhật dữ liệu mà không tải lại trang.",
    6: "Phần báo cáo cá nhân tập trung vào ba Use Case tiêu biểu: thêm mới hồ sơ nhân viên, xem danh sách và tìm kiếm hồ sơ, điều chuyển phòng ban và bổ nhiệm chức vụ mới.",
    7: "Phân hệ này quản lý dữ liệu nhân sự cốt lõi và cơ cấu tổ chức. Dữ liệu hồ sơ phải nhất quán với phòng ban, chức vụ và lịch sử điều chuyển để phục vụ tra cứu, hợp đồng, đánh giá và báo cáo.",
    8: "2. XÁC ĐỊNH MODULE/PHÂN HỆ ĐƯỢC PHÂN CÔNG",
    14: "3. THIẾT KẾ USE CASE",
    15: "Trong phạm vi cá nhân, lựa chọn 3 Use Case chính:",
    16: "UC-EMP-01 – Thêm mới hồ sơ nhân viên",
    17: "UC-EMP-02 – Xem danh sách, tìm kiếm và lọc hồ sơ nhân viên",
    18: "UC-ORG-04 – Điều chuyển phòng ban và bổ nhiệm chức vụ mới",
    19: "3.1. UC-EMP-01 – Thêm mới hồ sơ nhân viên",
    21: "Luồng chính:",
    22: "Chuyên viên HR chọn chức năng thêm mới nhân viên.",
    23: "Hệ thống hiển thị biểu mẫu hồ sơ nhân viên.",
    24: "HR nhập thông tin định danh, liên hệ và trạng thái làm việc.",
    25: "HR chọn phòng ban và chức vụ từ danh mục.",
    26: "Hệ thống kiểm tra dữ liệu bắt buộc, định dạng và tính duy nhất của CCCD, email.",
    27: "Nếu hợp lệ, hệ thống sinh mã nhân viên và lưu hồ sơ.",
    28: "Hệ thống thông báo thêm mới thành công.",
    29: "Luồng ngoại lệ:",
    30: "Thiếu hoặc sai định dạng dữ liệu: hiển thị lỗi tại trường nhập.",
    31: "CCCD hoặc email đã tồn tại: không lưu và yêu cầu nhập lại.",
    32: "Phòng ban hoặc chức vụ không hợp lệ: yêu cầu chọn lại danh mục.",
    33: "3.2. UC-EMP-02 – Xem danh sách, tìm kiếm và lọc hồ sơ nhân viên",
    36: "Luồng chính:",
    37: "HR hoặc Trưởng phòng mở danh sách hồ sơ nhân viên.",
    38: "Hệ thống tải trang dữ liệu đầu tiên theo quyền truy cập.",
    39: "Người dùng nhập từ khóa hoặc chọn phòng ban, trạng thái làm việc.",
    40: "Ajax gửi điều kiện tìm kiếm và số trang đến Controller.",
    41: "Hệ thống truy vấn dữ liệu, phân trang và trả về danh sách phù hợp.",
    42: "Giao diện cập nhật bảng kết quả mà không tải lại toàn trang.",
    43: "Luồng ngoại lệ:",
    44: "Không có dữ liệu phù hợp: hiển thị danh sách rỗng và thông báo.",
    45: "Điều kiện lọc không hợp lệ: bỏ qua giá trị sai và giữ bộ lọc hợp lệ.",
    46: "Yêu cầu Ajax lỗi: hiển thị thông báo và cho phép tải lại dữ liệu.",
    47: "3.3. UC-ORG-04 – Điều chuyển phòng ban và bổ nhiệm chức vụ mới",
    49: "Luồng chính:",
    50: "HR tìm và mở hồ sơ nhân viên cần điều chuyển.",
    51: "Hệ thống hiển thị phòng ban và chức vụ hiện tại.",
    52: "HR chọn phòng ban, chức vụ mới, ngày hiệu lực và nhập lý do.",
    53: "Hệ thống kiểm tra trạng thái nhân viên và ràng buộc Trưởng phòng.",
    54: "Hệ thống cập nhật hồ sơ và ghi lịch sử điều chuyển.",
    55: "Hệ thống thông báo điều chuyển thành công.",
    56: "Luồng ngoại lệ:",
    57: "Phòng ban mới trùng phòng ban hiện tại: cảnh báo và yêu cầu chọn lại.",
    58: "Nhân viên là Trưởng phòng cũ nhưng chưa có người thay thế: từ chối điều chuyển.",
    59: "4. THIẾT KẾ DỮ LIỆU",
    61: "Quan hệ chính: Một phòng ban có nhiều nhân viên; mỗi nhân viên tại một thời điểm thuộc một phòng ban và một chức vụ; một nhân viên có nhiều bản ghi lịch sử điều chuyển.",
    62: "4.1. ViewModel/EditModel",
    64: "5. THIẾT KẾ MVC CRUD FLOW",
    65: "5.1. Thêm mới hồ sơ nhân viên",
    67: "5.2. Tìm kiếm và lọc hồ sơ nhân viên",
    69: "5.3. Điều chuyển phòng ban và bổ nhiệm chức vụ",
    71: "6. KIỂM THỬ CÁC CHỨC NĂNG ĐƯỢC PHÂN CÔNG",
    73: "7. KẾT QUẢ THỰC HIỆN VÀ ĐÁNH GIÁ",
    74: "Phần cá nhân đã xác định và thiết kế phân hệ Phòng ban và Hồ sơ nhân sự. Ba Use Case được mô tả theo tác nhân, điều kiện, luồng chính và ngoại lệ; dữ liệu, ViewModel và luồng MVC được xác định làm cơ sở triển khai. Thiết kế bảo đảm tìm kiếm và phân trang bằng Ajax, kiểm tra trùng dữ liệu và lưu lịch sử điều chuyển.",
    75: "8. KẾT LUẬN",
    76: "Qua quá trình thực hiện phần cá nhân, em đã phân tích và thiết kế các chức năng cốt lõi của phân hệ Phòng ban và Hồ sơ nhân sự. Phân hệ giúp tập trung hóa hồ sơ, tra cứu nhanh, duy trì cơ cấu tổ chức chính xác và bảo toàn lịch sử thay đổi của nhân viên.",
}

TABLES = [
    [
        ["Mã", "Chức năng", "Tác nhân", "Nội dung chính"],
        ["FR-01", "Quản lý phòng ban", "HR, Admin", "Thêm, sửa, xóa và tra cứu phòng ban"],
        ["FR-02", "Quản lý chức vụ", "HR, Admin", "Thêm, sửa, xóa danh mục chức vụ"],
        ["FR-03", "Quản lý hồ sơ", "HR", "Thêm, cập nhật và xóa mềm hồ sơ"],
        ["FR-04", "Tìm kiếm hồ sơ", "HR, Trưởng phòng", "Lọc và phân trang bằng Ajax"],
        ["FR-05", "Điều chuyển nhân sự", "HR", "Cập nhật cơ cấu và lưu lịch sử"],
    ],
    [
        ["Nội dung", "Mô tả"],
        ["Mã Use Case", "UC-EMP-01"],
        ["Tên Use Case", "Thêm mới hồ sơ nhân viên"],
        ["Tác nhân", "Chuyên viên HR"],
        ["Mục đích", "Tạo hồ sơ nhân sự đầy đủ và gắn với cơ cấu tổ chức"],
        ["Điều kiện trước", "HR đã đăng nhập; phòng ban và chức vụ đã tồn tại"],
        ["Điều kiện sau", "Hồ sơ mới được lưu với mã nhân viên duy nhất"],
    ],
    [
        ["Nội dung", "Mô tả"],
        ["Mã Use Case", "UC-EMP-02"],
        ["Tên Use Case", "Xem danh sách, tìm kiếm và lọc hồ sơ nhân viên"],
        ["Tác nhân", "Chuyên viên HR, Trưởng phòng"],
        ["Mục đích", "Tra cứu hồ sơ nhanh theo từ khóa, phòng ban và trạng thái"],
        ["Điều kiện trước", "Người dùng đã đăng nhập và có quyền xem hồ sơ"],
        ["Điều kiện sau", "Danh sách phù hợp được hiển thị và phân trang bằng Ajax"],
    ],
    [
        ["Nội dung", "Mô tả"],
        ["Mã Use Case", "UC-ORG-04"],
        ["Tên Use Case", "Điều chuyển phòng ban và bổ nhiệm chức vụ mới"],
        ["Tác nhân", "Chuyên viên HR"],
        ["Mục đích", "Cập nhật cơ cấu công tác và lưu vết thay đổi"],
        ["Điều kiện trước", "Nhân viên đang làm việc; phòng ban và chức vụ mới hợp lệ"],
        ["Điều kiện sau", "Hồ sơ được cập nhật và lịch sử điều chuyển được lưu"],
    ],
    [
        ["Entity", "Thuộc tính chính", "Vai trò"],
        ["Employee", "EmployeeId, FullName, CitizenId, Email, DepartmentId, PositionId, Status", "Lưu hồ sơ và vị trí công tác hiện tại"],
        ["Department", "DepartmentId, DepartmentName, Description, ManagerId", "Lưu cơ cấu phòng ban và Trưởng phòng"],
        ["Position", "PositionId, PositionName, Description", "Lưu danh mục chức vụ"],
        ["TransferHistory", "TransferId, EmployeeId, PreviousDepartmentId, NewDepartmentId, PreviousPositionId, NewPositionId, EffectiveDate", "Lưu vết điều chuyển và bổ nhiệm"],
    ],
    [
        ["Model", "Dữ liệu chính", "Mục đích"],
        ["EmployeeCreateViewModel", "FullName, CitizenId, Email, DepartmentId, PositionId, Status", "Thêm hồ sơ"],
        ["EmployeeFilterViewModel", "Keyword, DepartmentId, Status, PageNumber", "Tìm kiếm và phân trang"],
        ["EmployeeTransferViewModel", "EmployeeId, NewDepartmentId, NewPositionId, EffectiveDate, Reason", "Điều chuyển"],
        ["DepartmentViewModel", "DepartmentId, DepartmentName, ManagerId", "Hiển thị cơ cấu phòng ban"],
    ],
    [
        ["Thành phần", "Luồng xử lý"],
        ["View", "Hiển thị biểu mẫu thêm hồ sơ"],
        ["Controller", "Nhận dữ liệu POST và kiểm tra ModelState"],
        ["Model/Service", "Kiểm tra trùng CCCD, email và danh mục liên quan"],
        ["Database", "Lưu bản ghi NhanVien"],
        ["Kết quả", "Chuyển đến chi tiết hồ sơ hoặc trả lỗi hợp lệ"],
    ],
    [
        ["Thành phần", "Luồng xử lý"],
        ["View", "Gửi từ khóa, bộ lọc và số trang bằng Ajax"],
        ["Controller", "Tiếp nhận tham số và kiểm tra quyền truy cập"],
        ["Model/Service", "Xây dựng truy vấn lọc và phân trang"],
        ["Database", "Truy vấn hồ sơ cùng phòng ban, chức vụ"],
        ["Kết quả", "Trả Partial View hoặc JSON để cập nhật bảng"],
    ],
    [
        ["Thành phần", "Luồng xử lý"],
        ["View", "Hiển thị thông tin hiện tại và biểu mẫu điều chuyển"],
        ["Controller", "Nhận yêu cầu và kiểm tra dữ liệu"],
        ["Model/Service", "Kiểm tra ràng buộc Trưởng phòng và tạo lịch sử"],
        ["Database", "Cập nhật NhanVien, thêm LichSuDieuChuyen trong giao dịch"],
        ["Kết quả", "Thông báo thành công hoặc lỗi nghiệp vụ"],
    ],
    [
        ["Mã", "Chức năng /", "Thao tác kiểm thử"],
        ["TC-01", "Thêm hồ sơ", "Nhập dữ liệu hợp lệ"],
        ["TC-02", "Trùng CCCD", "Dùng CCCD đã tồn tại"],
        ["TC-03", "Tìm kiếm", "Nhập mã hoặc tên nhân viên"],
        ["TC-04", "Lọc Ajax", "Chọn phòng ban và trạng thái"],
        ["TC-05", "Phân trang", "Chuyển trang không tải lại toàn trang"],
        ["TC-06", "Điều chuyển", "Chọn phòng ban, chức vụ mới hợp lệ"],
        ["TC-07", "Trùng phòng ban", "Chọn lại phòng ban hiện tại"],
        ["TC-08", "Ràng buộc Trưởng phòng", "Điều chuyển khi chưa chỉ định người thay thế"],
    ],
]


def replace_paragraph(paragraph, text):
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._element.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._element.rPr)
    for child in list(paragraph._element):
        if child.tag != qn("w:pPr"):
            paragraph._element.remove(child)
    run = paragraph.add_run()
    if run_properties is not None:
        run._element.insert(0, run_properties)
    parts = text.split("\n")
    for index, part in enumerate(parts):
        if index:
            run.add_break()
        run.add_text(part)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def main(reference, output):
    doc = Document(reference)
    if len(doc.tables) != len(TABLES):
        raise ValueError("Cấu trúc bảng của mẫu không khớp")
    for index, text in PARAGRAPHS.items():
        replace_paragraph(doc.paragraphs[index - 1], text)
    for table, rows in zip(doc.tables, TABLES):
        if len(table.rows) != len(rows):
            raise ValueError("Số hàng của bảng mẫu không khớp")
        set_repeat_table_header(table.rows[0])
        for row, values in zip(table.rows, rows):
            if len(row.cells) != len(values):
                raise ValueError("Số cột của bảng mẫu không khớp")
            for cell, value in zip(row.cells, values):
                replace_paragraph(cell.paragraphs[0], value)
    settings = doc.settings._element
    if settings.find(qn("w:updateFields")) is None:
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        settings.append(update_fields)
    doc.core_properties.title = "Báo cáo cá nhân thiết kế phân hệ phòng ban và hồ sơ nhân sự"
    doc.core_properties.author = "Vũ Mạnh Tiến"
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    main(argv[1], argv[2])
