from pathlib import Path
from sys import argv

from docx import Document


def iter_table(table, prefix):
    for row_index, row in enumerate(table.rows, 1):
        values = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
        print(f"{prefix} ROW {row_index}: " + " | ".join(values))


def main(path):
    doc = Document(path)
    print(f"FILE: {path}")
    print(f"SECTIONS: {len(doc.sections)}")
    for index, paragraph in enumerate(doc.paragraphs, 1):
        text = paragraph.text.strip()
        if text:
            print(f"P{index} [{paragraph.style.name}]: {text}")
    for index, table in enumerate(doc.tables, 1):
        print(f"TABLE {index}")
        iter_table(table, f"T{index}")
    for section_index, section in enumerate(doc.sections, 1):
        for name, container in (("HEADER", section.header), ("FOOTER", section.footer)):
            for paragraph_index, paragraph in enumerate(container.paragraphs, 1):
                text = paragraph.text.strip()
                if text:
                    print(f"S{section_index} {name} P{paragraph_index}: {text}")
            for table_index, table in enumerate(container.tables, 1):
                iter_table(table, f"S{section_index} {name} T{table_index}")


if __name__ == "__main__":
    main(Path(argv[1]))
