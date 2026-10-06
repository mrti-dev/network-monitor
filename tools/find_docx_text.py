from pathlib import Path
from sys import argv

from docx import Document


def main(path, query):
    doc = Document(path)
    items = []
    items.extend((f"P{i}", p.text) for i, p in enumerate(doc.paragraphs, 1))
    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            items.append((f"T{ti}R{ri}", " | ".join(cell.text for cell in row.cells)))
    for index, (label, text) in enumerate(items):
        if query.casefold() in text.casefold():
            for near_label, near_text in items[max(0, index - 3):index + 4]:
                print(f"{near_label}: {near_text}")
            print("---")


if __name__ == "__main__":
    main(Path(argv[1]), argv[2])
