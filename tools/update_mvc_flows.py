from copy import deepcopy
from pathlib import Path
from sys import argv

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


FLOWS = [
    [
        ["Thành phần", "Luồng xử lý chi tiết"],
        ["1. Client Request GET", "HR mở trang Employees/Create. Trình duyệt gửi HTTP GET đến EmployeeController.Create()."],
        ["2. Controller GET", "Kiểm tra quyền HR; gọi DepartmentService và PositionService lấy danh mục đang hoạt động; khởi tạo EmployeeCreateViewModel."],
        ["3. View", "Create.cshtml nhận ViewModel, hiển thị form thông tin cá nhân, Department, Position và WorkStatus; kèm anti-forgery token và validation message."],
        ["4. Client Request POST", "HR nhập dữ liệu và nhấn Lưu. Trình duyệt gửi HTTP POST đến EmployeeController.Create(EmployeeCreateViewModel model)."],
        ["5. Validation", "Controller kiểm tra ModelState; Service kiểm tra tuổi hợp lệ, IdentityCard và Email duy nhất, DepartmentId và PositionId còn tồn tại."],
        ["6. Mapping", "Nếu hợp lệ, EmployeeService sinh EmployeeCode, ánh xạ ViewModel sang Entity Employee và gán WorkStatus, HireDate, IsDeleted = false."],
        ["7. Data Access", "Repository thêm Employee; Entity Framework phát sinh INSERT và SaveChangesAsync() trong transaction."],
        ["8. Response", "Thành công: RedirectToAction(Details) kèm thông báo. Thất bại: rollback, nạp lại danh mục và trả View cùng lỗi đúng trường."],
    ],
    [
        ["Thành phần", "Luồng xử lý chi tiết"],
        ["1. Client Request", "HR hoặc Trưởng phòng mở Employees/Index. Trình duyệt gửi Keyword, DepartmentId, WorkStatus, PageNumber và PageSize."],
        ["2. Controller", "EmployeeController.Index(EmployeeFilterViewModel filter) kiểm tra quyền; Trưởng phòng chỉ được dùng DepartmentId của phòng mình quản lý."],
        ["3. Query Service", "EmployeeService tạo IQueryable<Employee>, loại bản ghi IsDeleted, áp dụng tìm kiếm theo EmployeeCode hoặc FullName và các bộ lọc hợp lệ."],
        ["4. Data Access", "Entity Framework dùng AsNoTracking(), Include Department/Position, CountAsync() tổng số bản ghi, sau đó Skip/Take để phân trang."],
        ["5. ViewModel", "Service ánh xạ kết quả sang danh sách EmployeeListItemViewModel và trả lại EmployeeFilterViewModel cùng TotalItems, TotalPages."],
        ["6. Full-page Response", "Yêu cầu GET thông thường trả View Index.cshtml chứa bộ lọc, bảng dữ liệu và điều khiển phân trang."],
        ["7. Ajax Request", "Khi thay đổi từ khóa, bộ lọc hoặc trang, JavaScript hủy request cũ nếu có và gửi GET Ajax với toàn bộ trạng thái lọc."],
        ["8. Partial Response", "Controller trả PartialView _EmployeeTable; JavaScript thay vùng bảng và phân trang, đồng thời cập nhật URL bằng history.replaceState()."],
        ["9. Exception Flow", "Không có dữ liệu: hiển thị trạng thái rỗng. Tham số sai: chuẩn hóa về mặc định. Lỗi mạng: giữ bảng cũ và cho phép thử lại."],
    ],
    [
        ["Thành phần", "Luồng xử lý chi tiết"],
        ["1. Client Request GET", "HR chọn Điều chuyển tại hồ sơ nhân viên. Trình duyệt gửi EmployeeId đến EmployeeController.Transfer(int id)."],
        ["2. Controller GET", "Kiểm tra quyền HR; tải Employee cùng Department, Position hiện tại; trả NotFound nếu hồ sơ không tồn tại hoặc đã nghỉ việc."],
        ["3. ViewModel", "Khởi tạo JobTransferEditModel gồm EmployeeId, NewDepartmentId, NewPositionId, TransferDate và Reason; nạp danh mục đích hợp lệ."],
        ["4. View", "Transfer.cshtml hiển thị cơ cấu hiện tại, danh mục mới, ngày hiệu lực, lý do; gửi anti-forgery token trong HTTP POST."],
        ["5. Controller POST", "EmployeeController.Transfer(JobTransferEditModel model) kiểm tra ModelState và gọi JobTransferService.TransferAsync(model, currentUserId)."],
        ["6. Business Rules", "Kiểm tra Department/Position mới tồn tại; phòng ban mới khác hiện tại; TransferDate hợp lệ; nếu nhân viên là Trưởng phòng thì phải có người thay thế."],
        ["7. Transaction", "Mở transaction; lưu DepartmentId và PositionId cũ; cập nhật Employee; tạo JobTransferHistory chứa giá trị cũ, mới, TransferDate và Reason."],
        ["8. Data Access", "Entity Framework phát sinh UPDATE Employee và INSERT JobTransferHistory; SaveChangesAsync(), sau đó commit. Có lỗi thì rollback."],
        ["9. Response", "Thành công: RedirectToAction(Details) kèm thông báo. Vi phạm nghiệp vụ: trả View với lỗi. Xung đột đồng thời: yêu cầu tải lại dữ liệu."],
    ],
]


def replace_paragraph(paragraph, text):
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._element.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._element.rPr)
    for child in list(paragraph._element):
        if child.tag != qn("w:pPr"):
            paragraph._element.remove(child)
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._element.insert(0, run_properties)


def set_table_rows(table, content):
    template_row = deepcopy(table.rows[-1]._tr)
    while len(table.rows) < len(content):
        table._tbl.append(deepcopy(template_row))
    while len(table.rows) > len(content):
        table._tbl.remove(table.rows[-1]._tr)
    for row_index, (row, values) in enumerate(zip(table.rows, content)):
        row.height = None
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell, value in zip(row.cells, values):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            replace_paragraph(cell.paragraphs[0], value)
        if row_index == 0:
            if tr_pr.find(qn("w:tblHeader")) is None:
                header = OxmlElement("w:tblHeader")
                header.set(qn("w:val"), "true")
                tr_pr.append(header)


def main(source, output):
    doc = Document(source)
    for table, content in zip(doc.tables[6:9], FLOWS):
        set_table_rows(table, content)
    for index in (64, 66, 68):
        doc.paragraphs[index].paragraph_format.keep_with_next = True
    doc.paragraphs[68].paragraph_format.page_break_before = True
    doc.core_properties.title = "Báo cáo cá nhân thiết kế phân hệ phòng ban và hồ sơ nhân sự"
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    main(argv[1], argv[2])
