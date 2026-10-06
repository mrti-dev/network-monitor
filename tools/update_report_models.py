from copy import deepcopy
from sys import argv

from docx import Document
from docx.oxml.ns import qn


ENTITY_ROWS = [
    ["Entity", "Key properties", "Role"],
    ["Employee", "EmployeeId, EmployeeCode, FullName, IdentityCard, Email, PhoneNumber, DepartmentId, PositionId, WorkStatus, HireDate, IsDeleted", "Lưu hồ sơ và vị trí công tác hiện tại"],
    ["Department", "DepartmentId, DepartmentCode, DepartmentName, ManagerId, PhoneNumber, IsDeleted", "Lưu cơ cấu phòng ban và Trưởng phòng"],
    ["Position", "PositionId, PositionName, BaseSalaryRate, Allowance, IsDeleted", "Lưu danh mục chức vụ"],
    ["JobTransferHistory", "TransferId, EmployeeId, OldDepartmentId, NewDepartmentId, OldPositionId, NewPositionId, TransferDate, Reason", "Lưu vết điều chuyển và bổ nhiệm"],
]

VIEW_MODEL_ROWS = [
    ["ViewModel/EditModel", "Key properties", "Purpose"],
    ["EmployeeCreateViewModel", "FullName, DateOfBirth, Gender, IdentityCard, Email, PhoneNumber, Address, AvatarUrl, DepartmentId, PositionId, WorkStatus, HireDate", "Thêm hồ sơ"],
    ["EmployeeFilterViewModel", "Keyword, DepartmentId, WorkStatus, PageNumber", "Tìm kiếm và phân trang"],
    ["JobTransferEditModel", "EmployeeId, NewDepartmentId, NewPositionId, TransferDate, Reason", "Điều chuyển"],
    ["DepartmentViewModel", "DepartmentId, DepartmentCode, DepartmentName, ManagerId, PhoneNumber", "Hiển thị cơ cấu phòng ban"],
]


def replace_paragraph(paragraph, text):
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._element.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._element.rPr)
    for child in list(paragraph._element):
        if child.tag != qn("w:pPr"):
            paragraph._element.remove(child)
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._element.insert(0, run_properties)


def replace_table(table, rows):
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            replace_paragraph(cell.paragraphs[0], value)


def main(source, output):
    doc = Document(source)
    replace_table(doc.tables[4], ENTITY_ROWS)
    replace_table(doc.tables[5], VIEW_MODEL_ROWS)
    doc.save(output)


if __name__ == "__main__":
    main(argv[1], argv[2])
