from pathlib import Path
from sys import argv

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


FLOWS = [
    (
        "Luồng MVC 1: Thêm mới hồ sơ nhân viên (UC-EMP-01)",
        "Chức năng này tiếp nhận hồ sơ nhân viên mới, kiểm tra dữ liệu định danh và bảo đảm phòng ban, chức vụ được gán còn hiệu lực trước khi lưu vào cơ sở dữ liệu.",
        [
            ("Client Request (User Action)", [
                "Chuyên viên HR chọn chức năng Thêm mới nhân viên tại màn hình quản lý hồ sơ.",
                "Trình duyệt gửi HTTP GET /Employees/Create đến máy chủ ứng dụng."
            ]),
            ("Controller (GET - Tải biểu mẫu)", [
                "EmployeeController tiếp nhận yêu cầu tại phương thức Create().",
                "Hệ thống kiểm tra quyền HR, gọi DepartmentService và PositionService để lấy các danh mục đang hoạt động, sau đó khởi tạo EmployeeCreateViewModel."
            ]),
            ("View (Hiển thị biểu mẫu)", [
                "Create.cshtml nhận ViewModel và hiển thị các trường thông tin cá nhân, phòng ban, chức vụ, trạng thái làm việc; đồng thời tạo anti-forgery token và vùng hiển thị lỗi kiểm tra dữ liệu."
            ]),
            ("Client Request (POST - Gửi dữ liệu)", [
                "Sau khi nhập dữ liệu, HR nhấn Lưu; trình duyệt gửi HTTP POST /Employees/Create kèm EmployeeCreateViewModel và anti-forgery token."
            ]),
            ("Controller & Validation (Kiểm tra dữ liệu)", [
                "EmployeeController.Create(EmployeeCreateViewModel model) kiểm tra ModelState.",
                "EmployeeService kiểm tra tuổi hợp lệ, tính duy nhất của IdentityCard và Email, đồng thời xác nhận DepartmentId và PositionId còn tồn tại."
            ]),
            ("Service & Mapping (Xử lý nghiệp vụ)", [
                "Nếu dữ liệu hợp lệ, EmployeeService sinh EmployeeCode, ánh xạ ViewModel sang Entity Employee và gán WorkStatus, HireDate, IsDeleted = false."
            ]),
            ("Data Access (Lưu cơ sở dữ liệu)", [
                "Repository thêm Employee; Entity Framework Core phát sinh lệnh INSERT và gọi SaveChangesAsync() trong transaction. Nếu xảy ra lỗi, transaction được rollback."
            ]),
            ("Response (Phản hồi kết quả)", [
                "Thành công: Controller chuyển đến trang Details và hiển thị thông báo tạo hồ sơ thành công.",
                "Thất bại: Controller nạp lại danh mục, trả về Create.cshtml cùng ViewModel và lỗi đúng tại trường nhập."
            ])
        ]
    ),
    (
        "Luồng MVC 2: Tìm kiếm và lọc hồ sơ nhân viên (UC-EMP-02)",
        "Chức năng này hỗ trợ HR và Trưởng phòng tra cứu hồ sơ theo từ khóa, phòng ban và trạng thái làm việc; kết quả được phân trang và cập nhật bằng Ajax mà không tải lại toàn bộ trang.",
        [
            ("Client Request (User Action)", [
                "Người dùng mở Employees/Index hoặc thay đổi Keyword, DepartmentId, WorkStatus, PageNumber và PageSize trên bộ lọc.",
                "JavaScript gửi HTTP GET; với thao tác lọc hoặc chuyển trang, yêu cầu được gửi dưới dạng Ajax và giữ nguyên toàn bộ trạng thái lọc."
            ]),
            ("Controller (Tiếp nhận & Kiểm soát phạm vi dữ liệu)", [
                "EmployeeController.Index(EmployeeFilterViewModel filter) kiểm tra quyền truy cập và chuẩn hóa tham số phân trang.",
                "Nếu người dùng là Trưởng phòng, Controller ghi đè DepartmentId bằng phòng ban được phân công để ngăn truy cập hồ sơ ngoài phạm vi quản lý."
            ]),
            ("Service (Xây dựng truy vấn)", [
                "EmployeeService tạo IQueryable<Employee>, loại các bản ghi IsDeleted và áp dụng tìm kiếm theo EmployeeCode hoặc FullName cùng các bộ lọc hợp lệ.",
                "Truy vấn sử dụng AsNoTracking(), Include Department và Position, CountAsync() để tính tổng số bản ghi, sau đó Skip/Take để phân trang."
            ]),
            ("ViewModel (Đóng gói kết quả)", [
                "Service ánh xạ kết quả sang EmployeeListItemViewModel và trả về EmployeeFilterViewModel chứa danh sách, TotalItems, TotalPages và trạng thái bộ lọc hiện hành."
            ]),
            ("View & Partial View (Hiển thị kết quả)", [
                "Yêu cầu tải trang thông thường trả về Index.cshtml gồm bộ lọc, bảng dữ liệu và điều khiển phân trang.",
                "Yêu cầu Ajax trả về PartialView _EmployeeTable; JavaScript thay vùng bảng, phân trang và cập nhật URL bằng history.replaceState()."
            ]),
            ("Exception Flow (Xử lý ngoại lệ)", [
                "Không có dữ liệu phù hợp: View hiển thị trạng thái rỗng nhưng giữ nguyên điều kiện lọc.",
                "Tham số sai được chuẩn hóa về giá trị mặc định; nếu yêu cầu Ajax lỗi, giao diện giữ bảng hiện tại, hiển thị thông báo và cho phép thử lại."
            ])
        ]
    ),
    (
        "Luồng MVC 3: Điều chuyển phòng ban và bổ nhiệm chức vụ (UC-ORG-04)",
        "Chức năng này cập nhật vị trí công tác hiện tại của nhân viên và đồng thời lưu lịch sử điều chuyển, bảo đảm dữ liệu hồ sơ và cơ cấu tổ chức nhất quán trong cùng một giao dịch.",
        [
            ("Client Request (GET - Mở biểu mẫu)", [
                "HR chọn Điều chuyển tại hồ sơ nhân viên; trình duyệt gửi EmployeeId đến HTTP GET /Employees/Transfer/{id}."
            ]),
            ("Controller (GET - Tải dữ liệu hiện tại)", [
                "EmployeeController.Transfer(int id) kiểm tra quyền HR và tải Employee cùng Department, Position hiện tại.",
                "Nếu hồ sơ không tồn tại, đã bị xóa mềm hoặc đã nghỉ việc, Controller trả về NotFound hoặc thông báo nghiệp vụ phù hợp."
            ]),
            ("ViewModel & View (Khởi tạo biểu mẫu)", [
                "Controller tạo JobTransferEditModel gồm EmployeeId, NewDepartmentId, NewPositionId, TransferDate và Reason; đồng thời nạp danh mục đích hợp lệ.",
                "Transfer.cshtml hiển thị cơ cấu hiện tại, danh mục mới, ngày hiệu lực và lý do; biểu mẫu gửi kèm anti-forgery token."
            ]),
            ("Controller (POST - Tiếp nhận yêu cầu)", [
                "EmployeeController.Transfer(JobTransferEditModel model) kiểm tra ModelState và gọi JobTransferService.TransferAsync(model, currentUserId)."
            ]),
            ("Service (Kiểm tra quy tắc nghiệp vụ)", [
                "Service xác nhận Department và Position mới tồn tại, phòng ban mới khác phòng ban hiện tại, TransferDate hợp lệ và nhân viên còn làm việc.",
                "Nếu nhân viên đang là Trưởng phòng, việc điều chuyển chỉ được tiếp tục khi đã có người thay thế hợp lệ."
            ]),
            ("Transaction & Data Access (Cập nhật đồng bộ)", [
                "Service mở transaction, lưu DepartmentId và PositionId cũ, cập nhật Entity Employee, sau đó tạo JobTransferHistory chứa giá trị cũ, giá trị mới, TransferDate và Reason.",
                "Entity Framework Core phát sinh UPDATE Employee và INSERT JobTransferHistory, gọi SaveChangesAsync() rồi commit; mọi lỗi làm transaction rollback."
            ]),
            ("Response (Phản hồi kết quả)", [
                "Thành công: Controller chuyển đến trang Details và hiển thị thông báo điều chuyển thành công.",
                "Vi phạm nghiệp vụ: trả lại Transfer.cshtml cùng thông báo lỗi; xung đột cập nhật đồng thời: yêu cầu tải lại dữ liệu trước khi thực hiện lại."
            ])
        ]
    )
]


def font(run, size=13, bold=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def body_paragraph(doc, text, first_line=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(0.5 if first_line else 0)
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(text))
    return p


def component(doc, number, title, items):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.keep_with_next = True
    font(p.add_run(f"{number}.  {title}:"), bold=True)
    for item in items:
        b = doc.add_paragraph()
        b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        b.paragraph_format.left_indent = Inches(1.0)
        b.paragraph_format.first_line_indent = Inches(-0.25)
        b.paragraph_format.space_after = Pt(2)
        font(b.add_run("•  "))
        font(b.add_run(item))


def main(source, output):
    doc = Document(source)
    start = doc.paragraphs[63]._element
    body = doc._element.body
    removing = False
    for child in list(body):
        if child is start:
            removing = True
        if removing and child.tag != qn("w:sectPr"):
            body.remove(child)

    title = doc.add_paragraph(style="Heading 2")
    title.paragraph_format.keep_with_next = True
    font(title.add_run("5. THIẾT KẾ KIẾN TRÚC LUỒNG MVC"), 17, True)
    body_paragraph(doc, "Phân hệ Quản lý phòng ban và hồ sơ nhân sự được triển khai theo mô hình ASP.NET Core MVC. Dưới đây là đặc tả chi tiết luồng xử lý của ba chức năng trọng tâm.")

    for flow_title, intro, steps in FLOWS:
        heading = doc.add_paragraph(style="Heading 3")
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(10)
        font(heading.add_run(flow_title), 13, True)
        body_paragraph(doc, intro)
        for number, (step_title, items) in enumerate(steps, 1):
            component(doc, number, step_title, items)

    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    main(argv[1], argv[2])
